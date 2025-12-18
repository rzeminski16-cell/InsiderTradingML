"""
Excel Report Generator for the Insider Trading Anomaly Detection System.

This module generates professional 8-sheet Excel reports with:
- Executive Summary
- Historical Anomalies
- Feature Engineering Details
- Model Performance & Comparison
- New Transaction Scoring
- Detailed Anomaly Explanations
- Technical Specifications
- Appendix & Methodology

Classes:
    ExcelReporter: Main class for report generation
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    NamedStyle,
    PatternFill,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from . import constants
from .utils import setup_logging, get_timestamp, format_number


class ExcelReporter:
    """
    Generate professional Excel reports for anomaly detection results.

    Creates an 8-sheet Excel workbook with comprehensive analysis,
    visualizations, and formatting.

    Attributes:
        output_filepath: Path to save the report
        workbook: openpyxl Workbook instance
        report_title: Title for the report
        report_description: Description of the report
    """

    def __init__(
        self,
        output_filepath: Optional[Union[str, Path]] = None,
        report_title: str = "Insider Trading Anomaly Detection Report",
        report_description: str = ""
    ):
        """
        Initialize the ExcelReporter.

        Args:
            output_filepath: Path to save report (auto-generated if None)
            report_title: Title for the report
            report_description: Description of the report
        """
        self.logger = setup_logging(__name__)

        if output_filepath is None:
            output_filepath = (
                constants.REPORTS_DIR /
                f"{get_timestamp()}_InsiderTrading_AnomalyReport.xlsx"
            )
        self.output_filepath = Path(output_filepath)
        self.output_filepath.parent.mkdir(parents=True, exist_ok=True)

        self.report_title = report_title
        self.report_description = report_description

        self.workbook = Workbook()
        self._setup_styles()

        # Report metadata
        self._metadata = {
            'generated_at': datetime.now(),
            'sheets_added': []
        }

    def _setup_styles(self) -> None:
        """Set up named styles for the workbook."""
        # Header style
        self.header_style = NamedStyle(name='header_style')
        self.header_style.font = Font(
            name='Arial',
            size=12,
            bold=True,
            color='FFFFFF'
        )
        self.header_style.fill = PatternFill(
            start_color='1E88E5',
            end_color='1E88E5',
            fill_type='solid'
        )
        self.header_style.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )
        self.header_style.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title style
        self.title_style = NamedStyle(name='title_style')
        self.title_style.font = Font(
            name='Arial',
            size=16,
            bold=True,
            color='333333'
        )
        self.title_style.alignment = Alignment(horizontal='left')

        # Data cell style
        self.data_style = NamedStyle(name='data_style')
        self.data_style.font = Font(name='Arial', size=11)
        self.data_style.border = Border(
            left=Side(style='thin', color='DEE2E6'),
            right=Side(style='thin', color='DEE2E6'),
            top=Side(style='thin', color='DEE2E6'),
            bottom=Side(style='thin', color='DEE2E6')
        )

        # Register styles
        try:
            self.workbook.add_named_style(self.header_style)
            self.workbook.add_named_style(self.title_style)
            self.workbook.add_named_style(self.data_style)
        except ValueError:
            pass  # Styles already exist

    def _get_fill_color(self, risk_level: str) -> PatternFill:
        """Get fill color based on risk level."""
        colors = {
            'CRITICAL': 'DC3545',
            'HIGH': 'FF4500',
            'MEDIUM': 'FFC107',
            'LOW': 'FFD700',
            'NORMAL': '28A745'
        }
        color = colors.get(risk_level, 'FFFFFF')
        return PatternFill(start_color=color, end_color=color, fill_type='solid')

    def _write_dataframe(
        self,
        ws,
        df: pd.DataFrame,
        start_row: int = 1,
        start_col: int = 1,
        include_header: bool = True,
        include_index: bool = False
    ) -> int:
        """
        Write a DataFrame to a worksheet with formatting.

        Returns the next available row after the DataFrame.
        """
        rows = list(dataframe_to_rows(df, index=include_index, header=include_header))

        for r_idx, row in enumerate(rows, start=start_row):
            for c_idx, value in enumerate(row, start=start_col):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                if r_idx == start_row and include_header:
                    cell.style = self.header_style
                else:
                    cell.style = self.data_style

        # Auto-fit column widths
        for col_idx in range(start_col, start_col + len(df.columns)):
            col_letter = get_column_letter(col_idx)
            max_length = 0
            for cell in ws[col_letter]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

        return start_row + len(rows)

    def add_executive_summary(
        self,
        summary_data: Dict
    ) -> None:
        """
        Add Sheet 1: Executive Summary.

        Args:
            summary_data: Dictionary with summary statistics
        """
        ws = self.workbook.active
        ws.title = "Executive Summary"

        row = 1

        # Title
        ws.cell(row=row, column=1, value=self.report_title)
        ws.cell(row=row, column=1).style = self.title_style
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 2

        # Report date
        ws.cell(row=row, column=1, value="Report Generated:")
        ws.cell(row=row, column=2, value=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        row += 1

        # Data period
        if 'date_range' in summary_data:
            ws.cell(row=row, column=1, value="Data Period:")
            ws.cell(row=row, column=2,
                   value=f"{summary_data['date_range'].get('start', 'N/A')} to "
                         f"{summary_data['date_range'].get('end', 'N/A')}")
        row += 2

        # Key Metrics Section
        ws.cell(row=row, column=1, value="KEY METRICS")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 1

        metrics = [
            ("Total Transactions Analyzed", summary_data.get('total_transactions', 0)),
            ("Historical Anomalies Detected", summary_data.get('n_anomalies', 0)),
            ("Anomaly Rate", f"{summary_data.get('anomaly_rate', 0)*100:.1f}%"),
            ("New Transactions Scored", summary_data.get('n_new_transactions', 0)),
            ("Critical Risk Transactions", summary_data.get('n_critical', 0)),
            ("Models Used", summary_data.get('n_models', 0))
        ]

        for metric, value in metrics:
            ws.cell(row=row, column=1, value=metric)
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 1

        # Methodology Overview
        ws.cell(row=row, column=1, value="METHODOLOGY OVERVIEW")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 1

        methodology = [
            ("Aggregation Window", f"{summary_data.get('window_days', 14)} days"),
            ("Features Engineered", summary_data.get('n_features', 0)),
            ("Anomaly Threshold", summary_data.get('threshold_method', 'Percentile-based'))
        ]

        for item, value in methodology:
            ws.cell(row=row, column=1, value=item)
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 1

        # Key Findings
        if 'key_findings' in summary_data:
            ws.cell(row=row, column=1, value="KEY FINDINGS")
            ws.cell(row=row, column=1).font = Font(bold=True, size=14)
            row += 1

            for finding in summary_data['key_findings']:
                ws.cell(row=row, column=1, value=f"• {finding}")
                row += 1

        row += 1

        # Recommendations
        if 'recommendations' in summary_data:
            ws.cell(row=row, column=1, value="RECOMMENDATIONS")
            ws.cell(row=row, column=1).font = Font(bold=True, size=14)
            row += 1

            for i, rec in enumerate(summary_data['recommendations'], 1):
                ws.cell(row=row, column=1, value=f"Priority {i}: {rec}")
                row += 1

        self._metadata['sheets_added'].append('Executive Summary')
        self.logger.info("Added Executive Summary sheet")

    def add_historical_anomalies(
        self,
        anomalies_df: pd.DataFrame,
        models_list: List[str]
    ) -> None:
        """
        Add Sheet 2: Historical Anomalies.

        Args:
            anomalies_df: DataFrame with anomaly data
            models_list: List of model names used
        """
        ws = self.workbook.create_sheet("Historical Anomalies")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Historical Anomalies")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        # Write anomaly table
        if len(anomalies_df) > 0:
            row = self._write_dataframe(ws, anomalies_df, start_row=row)

            # Add conditional formatting for risk levels
            if 'risk_level' in anomalies_df.columns:
                risk_col_idx = list(anomalies_df.columns).index('risk_level') + 1
                for r in range(4, row):
                    cell = ws.cell(row=r, column=risk_col_idx)
                    risk = cell.value
                    if risk:
                        cell.fill = self._get_fill_color(risk)

            # Add color scale for anomaly scores
            if 'anomaly_score' in anomalies_df.columns:
                score_col_idx = list(anomalies_df.columns).index('anomaly_score') + 1
                score_col_letter = get_column_letter(score_col_idx)
                ws.conditional_formatting.add(
                    f'{score_col_letter}4:{score_col_letter}{row-1}',
                    ColorScaleRule(
                        start_type='min', start_color='28A745',
                        mid_type='percentile', mid_value=50, mid_color='FFC107',
                        end_type='max', end_color='DC3545'
                    )
                )

            row += 2

            # Add summary statistics
            ws.cell(row=row, column=1, value="Summary Statistics")
            ws.cell(row=row, column=1).font = Font(bold=True)
            row += 1

            ws.cell(row=row, column=1, value="Total Anomaly Windows:")
            ws.cell(row=row, column=2, value=len(anomalies_df[anomalies_df.get('is_anomaly', True) == 1]) if 'is_anomaly' in anomalies_df.columns else len(anomalies_df))
            row += 1

            if 'anomaly_score' in anomalies_df.columns:
                ws.cell(row=row, column=1, value="Average Anomaly Score:")
                ws.cell(row=row, column=2, value=f"{anomalies_df['anomaly_score'].mean():.3f}")
                row += 1

                ws.cell(row=row, column=1, value="Highest Score:")
                ws.cell(row=row, column=2, value=f"{anomalies_df['anomaly_score'].max():.3f}")

            # Add chart
            if 'anomaly_score' in anomalies_df.columns and len(anomalies_df) > 1:
                try:
                    chart = BarChart()
                    chart.title = "Anomaly Scores by Window"
                    chart.x_axis.title = "Window"
                    chart.y_axis.title = "Anomaly Score"

                    score_col = list(anomalies_df.columns).index('anomaly_score') + 1
                    data = Reference(ws, min_col=score_col, min_row=3,
                                   max_row=min(3 + len(anomalies_df), 23))
                    chart.add_data(data, titles_from_data=True)
                    chart.width = 15
                    chart.height = 8
                    ws.add_chart(chart, "G3")
                except Exception as e:
                    self.logger.warning(f"Could not add chart: {e}")

        else:
            ws.cell(row=row, column=1, value="No anomalies detected.")

        # Freeze header
        ws.freeze_panes = 'A4'

        # Enable auto filter
        if len(anomalies_df) > 0:
            ws.auto_filter.ref = f"A3:{get_column_letter(len(anomalies_df.columns))}{row-1}"

        self._metadata['sheets_added'].append('Historical Anomalies')
        self.logger.info("Added Historical Anomalies sheet")

    def add_feature_engineering_details(
        self,
        feature_stats: Dict[str, Dict],
        feature_importance: Optional[Dict[str, float]] = None,
        correlation_matrix: Optional[pd.DataFrame] = None
    ) -> None:
        """
        Add Sheet 3: Feature Engineering Details.

        Args:
            feature_stats: Dictionary of feature statistics
            feature_importance: Dictionary of feature importance scores
            correlation_matrix: Correlation matrix DataFrame
        """
        ws = self.workbook.create_sheet("Feature Engineering")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Feature Engineering Details")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        # Feature Statistics Table
        ws.cell(row=row, column=1, value="Feature Statistics")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        if feature_stats:
            stats_data = []
            for name, stats in feature_stats.items():
                stats_data.append({
                    'Feature Name': name,
                    'Mean': f"{stats.get('mean', 0):.4f}",
                    'Std Dev': f"{stats.get('std', 0):.4f}",
                    'Min': f"{stats.get('min', 0):.4f}",
                    'Max': f"{stats.get('max', 0):.4f}",
                    'Median': f"{stats.get('median', 0):.4f}"
                })

            stats_df = pd.DataFrame(stats_data)
            row = self._write_dataframe(ws, stats_df, start_row=row)

        row += 2

        # Feature Importance
        if feature_importance:
            ws.cell(row=row, column=1, value="Feature Importance")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            importance_data = [
                {'Feature': name, 'Importance Score': f"{score:.4f}"}
                for name, score in sorted(
                    feature_importance.items(),
                    key=lambda x: x[1],
                    reverse=True
                )
            ]
            importance_df = pd.DataFrame(importance_data)
            row = self._write_dataframe(ws, importance_df, start_row=row)

        row += 2

        # Feature Categories
        ws.cell(row=row, column=1, value="Feature Categories")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        for category, features in constants.FEATURE_CATEGORIES.items():
            ws.cell(row=row, column=1, value=category)
            ws.cell(row=row, column=1).font = Font(bold=True)
            row += 1
            for feature in features:
                ws.cell(row=row, column=1, value=f"  • {feature}")
                row += 1
            row += 1

        self._metadata['sheets_added'].append('Feature Engineering')
        self.logger.info("Added Feature Engineering Details sheet")

    def add_model_performance(
        self,
        training_metrics: Dict[str, Dict],
        model_comparison: Optional[Dict] = None
    ) -> None:
        """
        Add Sheet 4: Model Performance & Comparison.

        Args:
            training_metrics: Dictionary of training metrics per model
            model_comparison: Model comparison data
        """
        ws = self.workbook.create_sheet("Model Performance")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Model Performance & Comparison")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        # Model Training Summary
        ws.cell(row=row, column=1, value="Model Training Summary")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        if training_metrics:
            metrics_data = []
            for name, metrics in training_metrics.items():
                if 'error' not in metrics:
                    metrics_data.append({
                        'Model': name,
                        'Type': metrics.get('model_type', 'Unknown'),
                        'Training Time (s)': f"{metrics.get('training_time', 0):.2f}",
                        'Anomalies Detected': metrics.get('anomalies_detected', 0),
                        'Anomaly Rate': f"{metrics.get('anomaly_rate', 0)*100:.1f}%",
                        'Mean Score': f"{metrics.get('mean_score', 0):.3f}"
                    })

            if metrics_data:
                metrics_df = pd.DataFrame(metrics_data)
                row = self._write_dataframe(ws, metrics_df, start_row=row)

        row += 2

        # Model Comparison
        if model_comparison:
            ws.cell(row=row, column=1, value="Model Agreement Analysis")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            if 'agreement' in model_comparison:
                ws.cell(row=row, column=1, value="Full Agreement Rate:")
                ws.cell(row=row, column=2,
                       value=f"{model_comparison['agreement'].get('full_agreement', 0)*100:.1f}%")
                row += 1

            if 'unanimous_anomalies' in model_comparison:
                ws.cell(row=row, column=1, value="Unanimous Anomalies:")
                ws.cell(row=row, column=2,
                       value=len(model_comparison['unanimous_anomalies']))
                row += 1

            if 'contested_points' in model_comparison:
                ws.cell(row=row, column=1, value="Contested Points:")
                ws.cell(row=row, column=2,
                       value=len(model_comparison['contested_points']))

        self._metadata['sheets_added'].append('Model Performance')
        self.logger.info("Added Model Performance sheet")

    def add_new_transaction_scoring(
        self,
        new_transactions_df: pd.DataFrame,
        scores_summary: Optional[Dict] = None
    ) -> None:
        """
        Add Sheet 5: New Transaction Scoring.

        Args:
            new_transactions_df: DataFrame with scored new transactions
            scores_summary: Summary of scoring results
        """
        ws = self.workbook.create_sheet("New Transaction Scoring")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="New Transaction Scoring")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        if len(new_transactions_df) > 0:
            # Write transactions table
            row = self._write_dataframe(ws, new_transactions_df, start_row=row)

            # Apply conditional formatting
            if 'risk_level' in new_transactions_df.columns:
                risk_col_idx = list(new_transactions_df.columns).index('risk_level') + 1
                for r in range(4, row):
                    cell = ws.cell(row=r, column=risk_col_idx)
                    risk = cell.value
                    if risk:
                        cell.fill = self._get_fill_color(risk)

            row += 2

            # Summary
            ws.cell(row=row, column=1, value="Summary")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            if scores_summary:
                for key, value in scores_summary.items():
                    ws.cell(row=row, column=1, value=key)
                    ws.cell(row=row, column=2, value=str(value))
                    row += 1
            else:
                ws.cell(row=row, column=1, value="Total New Transactions:")
                ws.cell(row=row, column=2, value=len(new_transactions_df))
                row += 1

                if 'risk_level' in new_transactions_df.columns:
                    risk_counts = new_transactions_df['risk_level'].value_counts()
                    for risk, count in risk_counts.items():
                        ws.cell(row=row, column=1, value=f"{risk}:")
                        ws.cell(row=row, column=2, value=count)
                        row += 1
        else:
            ws.cell(row=row, column=1, value="No new transactions to score.")

        self._metadata['sheets_added'].append('New Transaction Scoring')
        self.logger.info("Added New Transaction Scoring sheet")

    def add_anomaly_explanations(
        self,
        explanations: List[Dict],
        feature_names: List[str]
    ) -> None:
        """
        Add Sheet 6: Detailed Anomaly Explanations.

        Args:
            explanations: List of anomaly explanation dictionaries
            feature_names: List of feature names
        """
        ws = self.workbook.create_sheet("Anomaly Explanations")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Detailed Anomaly Explanations")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        if not explanations:
            ws.cell(row=row, column=1, value="No anomaly explanations available.")
            self._metadata['sheets_added'].append('Anomaly Explanations')
            return

        for i, explanation in enumerate(explanations[:10], 1):  # Limit to top 10
            if 'error' in explanation:
                continue

            ws.cell(row=row, column=1, value=f"Anomaly #{i}")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            ws.cell(row=row, column=1, value="Anomaly Score:")
            ws.cell(row=row, column=2,
                   value=f"{explanation.get('anomaly_score', 0):.3f}")
            row += 1

            ws.cell(row=row, column=1, value="Significant Deviations:")
            ws.cell(row=row, column=2,
                   value=explanation.get('n_significant_deviations', 0))
            row += 1

            # Top reasons
            if 'top_reasons' in explanation and explanation['top_reasons']:
                ws.cell(row=row, column=1, value="Top Reasons:")
                row += 1
                for reason in explanation['top_reasons']:
                    ws.cell(row=row, column=1, value=f"  • {reason}")
                    row += 1

            row += 1

            # Feature deviations table
            if 'feature_deviations' in explanation:
                ws.cell(row=row, column=1, value="Feature Deviations:")
                row += 1

                dev_data = []
                for fname, dev in explanation['feature_deviations'].items():
                    if abs(dev.get('z_score', 0)) > 1:  # Only show significant
                        dev_data.append({
                            'Feature': fname,
                            'Value': f"{dev.get('value', 0):.3f}",
                            'Mean': f"{dev.get('mean', 0):.3f}",
                            'Z-Score': f"{dev.get('z_score', 0):.2f}"
                        })

                if dev_data:
                    dev_df = pd.DataFrame(dev_data)
                    row = self._write_dataframe(ws, dev_df, start_row=row)

            row += 2

        self._metadata['sheets_added'].append('Anomaly Explanations')
        self.logger.info("Added Anomaly Explanations sheet")

    def add_technical_specs(
        self,
        specs_data: Dict
    ) -> None:
        """
        Add Sheet 7: Technical Specifications.

        Args:
            specs_data: Dictionary with technical specifications
        """
        ws = self.workbook.create_sheet("Technical Specifications")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Technical Specifications")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        sections = [
            ("Data Specifications", [
                ("Raw Data Source", specs_data.get('data_source', 'N/A')),
                ("Date Range", specs_data.get('date_range', 'N/A')),
                ("Record Count", specs_data.get('n_transactions', 'N/A')),
                ("Window Count", specs_data.get('n_windows', 'N/A'))
            ]),
            ("Feature Engineering", [
                ("Aggregation Window", f"{specs_data.get('window_days', 14)} days"),
                ("Features Created", specs_data.get('n_features', 'N/A')),
                ("Standardization", "Z-score (mean=0, std=1)"),
                ("Missing Value Handling", "Forward fill, then drop")
            ]),
            ("Anomaly Threshold", [
                ("Method", specs_data.get('threshold_method', 'Percentile')),
                ("Parameter", specs_data.get('threshold_param', '95th percentile')),
                ("Threshold Value", f"{specs_data.get('threshold_value', 'N/A')}")
            ]),
            ("Report Generation", [
                ("Generated At", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
                ("Python Version", specs_data.get('python_version', 'N/A')),
                ("scikit-learn Version", specs_data.get('sklearn_version', 'N/A'))
            ])
        ]

        for section_title, items in sections:
            ws.cell(row=row, column=1, value=section_title)
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            for label, value in items:
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=str(value))
                row += 1

            row += 1

        # Model Specifications
        if 'models' in specs_data:
            ws.cell(row=row, column=1, value="Model Specifications")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1

            for model_name, model_specs in specs_data['models'].items():
                ws.cell(row=row, column=1, value=f"Model: {model_name}")
                ws.cell(row=row, column=1).font = Font(bold=True)
                row += 1

                for param, value in model_specs.items():
                    ws.cell(row=row, column=1, value=f"  {param}")
                    ws.cell(row=row, column=2, value=str(value))
                    row += 1

                row += 1

        self._metadata['sheets_added'].append('Technical Specifications')
        self.logger.info("Added Technical Specifications sheet")

    def add_appendix(
        self,
        methodology_notes: Optional[str] = None
    ) -> None:
        """
        Add Sheet 8: Appendix & Methodology.

        Args:
            methodology_notes: Additional methodology notes
        """
        ws = self.workbook.create_sheet("Appendix")

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Appendix - Glossary & Methodology")
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        # Glossary
        ws.cell(row=row, column=1, value="Term Definitions")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        glossary = [
            ("Anomaly Score", "Normalized [0-1] measure of how unusual a period is"),
            ("Anomaly Rate", "Percentage of windows/transactions flagged as anomalous"),
            ("Executive Concentration", "Herfindahl index measuring trading dominance"),
            ("Herfindahl Index", "Sum of squared market shares, measures concentration"),
            ("Risk Level", "Categorical classification: CRITICAL, HIGH, MEDIUM, LOW, NORMAL"),
            ("Window", "Time period over which transactions are aggregated"),
            ("Z-Score", "Number of standard deviations from the mean")
        ]

        for term, definition in glossary:
            ws.cell(row=row, column=1, value=term)
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=2, value=definition)
            row += 1

        row += 1

        # Methodology Overview
        ws.cell(row=row, column=1, value="Methodology Overview")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        methodology = [
            "1. Data Preparation: Raw transaction data is loaded, validated, and cleaned.",
            "2. Time Aggregation: Transactions are grouped into time windows.",
            "3. Feature Engineering: 20+ features are computed for each window.",
            "4. Model Training: Multiple unsupervised anomaly detection models are trained.",
            "5. Anomaly Detection: Historical anomalies are identified based on thresholds.",
            "6. Scoring: New transactions are scored using trained models.",
            "7. Reporting: Comprehensive reports are generated with visualizations."
        ]

        for step in methodology:
            ws.cell(row=row, column=1, value=step)
            row += 1

        row += 1

        # Limitations
        ws.cell(row=row, column=1, value="Limitations & Considerations")
        ws.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1

        limitations = [
            "• Model assumes historical data is representative of normal behavior",
            "• Cannot directly prove insider trading, only flags suspicious patterns",
            "• New market regimes may require model recalibration",
            "• Results should be reviewed by domain experts before action"
        ]

        for limitation in limitations:
            ws.cell(row=row, column=1, value=limitation)
            row += 1

        if methodology_notes:
            row += 1
            ws.cell(row=row, column=1, value="Additional Notes")
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            row += 1
            ws.cell(row=row, column=1, value=methodology_notes)

        self._metadata['sheets_added'].append('Appendix')
        self.logger.info("Added Appendix sheet")

    def save_and_validate(self) -> Path:
        """
        Save workbook and validate structure.

        Returns:
            Path to saved file
        """
        # Remove default sheet if empty
        if 'Sheet' in self.workbook.sheetnames:
            del self.workbook['Sheet']

        # Set document properties
        self.workbook.properties.title = self.report_title
        self.workbook.properties.creator = "Insider Trading Anomaly Detection System"
        self.workbook.properties.created = self._metadata['generated_at']

        # Save workbook
        self.workbook.save(self.output_filepath)

        self.logger.info(
            f"Report saved to {self.output_filepath} "
            f"({len(self._metadata['sheets_added'])} sheets)"
        )

        return self.output_filepath

    def generate_full_report(
        self,
        summary_data: Dict,
        anomalies_df: pd.DataFrame,
        feature_stats: Dict,
        training_metrics: Dict,
        new_transactions_df: Optional[pd.DataFrame] = None,
        explanations: Optional[List[Dict]] = None,
        specs_data: Optional[Dict] = None,
        feature_importance: Optional[Dict] = None,
        model_comparison: Optional[Dict] = None
    ) -> Path:
        """
        Generate complete report with all sheets.

        Args:
            summary_data: Executive summary data
            anomalies_df: Historical anomalies DataFrame
            feature_stats: Feature statistics
            training_metrics: Model training metrics
            new_transactions_df: Scored new transactions
            explanations: Anomaly explanations
            specs_data: Technical specifications
            feature_importance: Feature importance scores
            model_comparison: Model comparison data

        Returns:
            Path to saved report
        """
        # Sheet 1: Executive Summary
        self.add_executive_summary(summary_data)

        # Sheet 2: Historical Anomalies
        models_list = list(training_metrics.keys())
        self.add_historical_anomalies(anomalies_df, models_list)

        # Sheet 3: Feature Engineering
        self.add_feature_engineering_details(
            feature_stats,
            feature_importance
        )

        # Sheet 4: Model Performance
        self.add_model_performance(training_metrics, model_comparison)

        # Sheet 5: New Transaction Scoring
        if new_transactions_df is not None and len(new_transactions_df) > 0:
            self.add_new_transaction_scoring(new_transactions_df)
        else:
            self.add_new_transaction_scoring(pd.DataFrame())

        # Sheet 6: Anomaly Explanations
        feature_names = list(feature_stats.keys()) if feature_stats else []
        self.add_anomaly_explanations(explanations or [], feature_names)

        # Sheet 7: Technical Specifications
        self.add_technical_specs(specs_data or {})

        # Sheet 8: Appendix
        self.add_appendix()

        return self.save_and_validate()
